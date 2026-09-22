"""
_verify_report_ch13.py —— 核对第十三章「组员确认」

核对点：
1. 确认行已填写：成员姓名 / 电子签名 / 日期，日期为 2026.9.23
2. 签名与姓名在字体上有所区分（签名用楷体），便于与正文文字区分开
3. 表格保持模板的 4 行 3 列结构（单人组只填一行，其余留空）
4. 第十三章的确认用语与模板提示语未被改动
5. 全文 13 章齐全，结构未受破坏
"""

import sys
import zipfile

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

P = r"report\实验二_基于用户访谈的需求演化.docx"
PASS = FAIL = 0


def chk(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {name}" + (f"  {detail}" if detail else ""))
    else:
        FAIL += 1
        print(f"  [FAIL] {name}  {detail}")


d = Document(P)
t = d.tables[12]
rows = [[c.text.strip() for c in r.cells] for r in t.rows]
texts = [p.text.strip() for p in d.paragraphs]

print("【1】确认行填写")
chk("表头未变", rows[0] == ["成员姓名", "电子签名", "日期"], " | ".join(rows[0]))
chk("第 1 行姓名", rows[1][0] == "肖子恒", rows[1][0])
chk("第 1 行签名", rows[1][1] == "肖子恒", rows[1][1])
chk("第 1 行日期为 2026.9.23", rows[1][2] == "2026.9.23", rows[1][2])
left = [(ri, ci) for ri in (2, 3) for ci in range(3) if rows[ri][ci]]
chk("其余两行按模板留空（单人组）", not left, f"多填 {left}" if left else "")

print()
print("【2】签名与姓名在格式上区分")
c_name = t.rows[1].cells[0]
c_sign = t.rows[1].cells[1]
name_font = {r.font.name for p in c_name.paragraphs for r in p.runs}
sign_font = {r.font.name for p in c_sign.paragraphs for r in p.runs}
chk("签名使用楷体，姓名用默认字体",
    any("楷体" in (f or "") for f in sign_font) and not any("楷体" in (f or "") for f in name_font),
    f"姓名={name_font} 签名={sign_font}")

print()
print("【3】模板结构")
chk("表格仍为 4 行 3 列", len(t.rows) == 4 and len(t.columns) == 3,
    f"{len(t.rows)} 行 {len(t.columns)} 列")
chk("确认用语保留",
    "本人确认本报告所记录的访谈、需求分析、开发和测试过程与本组实际工作一致。" in texts)
chk("模板提示语保留", "单人组填写一行即可。" in texts)

print()
print("【4】全文完整性")
heads = [p.text.strip() for p in d.paragraphs if p.style.name.startswith("Heading")]
chk("13 个一级章节齐全", len([h for h in heads if h[:2] in
    ("一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、", "十、", "十一", "十二", "十三")]) == 13,
    " / ".join(h[:6] for h in heads))
chk("表格数仍为 13", len(d.tables) == 13, f"{len(d.tables)} 张")
chk("图 8-1 在位", len(d.inline_shapes) == 1)
z = zipfile.ZipFile(P)
chk("ZIP 结构完好", z.testzip() is None)

print()
print("【5】其余章节未被影响")
chk("第十一章 AI 记录表仍完整",
    all(d.tables[11].rows[r].cells[c].text.strip() for r in range(1, 5) for c in range(4)))
chk("第十二章 8 段回答仍在",
    sum(1 for i, x in enumerate(texts) if x.startswith(("报名环节由提交即生效", "数据与界面同步调整",
        "候补这条链路经三层追问", "审核与名额的先后", "最容易误解的是已报名人数",
        "另一处是候补的排序依据", "V1.0 的需求基本按自己", "另一变化是认识到单个角色"))) == 8)
chk("第五章管理员行仍是只读口径",
    "查看平台全部活动的报名与候补名单" in d.tables[6].rows[3].cells[1].text)

print()
print("=" * 62)
print(f"核验结果：{PASS} 项通过 / {FAIL} 项失败")
print("=" * 62)
sys.exit(0 if FAIL == 0 else 1)
